#!/usr/bin/env python3
"""Print all sheets, column headers, and first 5 rows of an Excel file.

Usage: python read_excel.py <path_to_xlsx>
"""
import sys

import openpyxl


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: read_excel.py <file.xlsx>")
        sys.exit(1)

    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n=== Sheet: {sheet_name!r} ({ws.max_row} rows × {ws.max_column} cols) ===")

        headers = [cell.value for cell in ws[1]]
        print("Headers:", headers)

        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if i > 6:  # first 5 data rows
                print(f"  ... ({ws.max_row - 1} total data rows)")
                break
            print(f"  Row {i}:", row)


if __name__ == "__main__":
    main()
